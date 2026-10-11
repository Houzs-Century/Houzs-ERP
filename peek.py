import openpyxl, sys
p=r"C:\Users\User\Downloads\houzs-century-ALL-SO-detail-with-CustomerPO-2026-09-11.xlsx"
wb=openpyxl.load_workbook(p, read_only=True, data_only=True)
for ws in wb.worksheets:
    print("SHEET:", ws.title, ws.max_row, ws.max_column)
    for i,row in enumerate(ws.iter_rows(values_only=True)):
        print(i, [str(c)[:40] if c is not None else "" for c in row])
        if i>=12: break
    print("---")
